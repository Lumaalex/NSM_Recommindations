# -*- coding: utf-8 -*-
"""
EXCEL -> JSON
=============
Читает мастер-базу «Банк_фраз.xlsx» и собирает из неё data/phrases.json.
После этого нужно выполнить  python tools/build.py  — HTML-конструктор пересоберётся
уже из ваших правок в Excel.

Можно одной командой:  python tools/build.py --from-excel

Формат листа «Фразы» (колонки):
  ID | Направление | Блок | Микроблок | Способность | Статус | Текст формулировки |
  Источник | Требует примера | Подсказка: какой пример нужен |
  Цели(ID) | Приёмы(ID) | Рекомендации(ID) | Теги/домен | Активна

Связи перечисляются через точку с запятой: GL-01;GL-07
Логические поля: «Да» / «Нет».

КОДЫ: с версии данных 1.2.0 коды микротем семантические — AS1sens…AS8vis (актуальное
состояние), RC1home…RC6peers (рекомендации); ID фраз вида AS5play-S2-07.
Старые коды (AS1-S3-01, микроблок AS4, RC1…) в Excel переживать не нужно: при сборке
они автоматически переводятся в новые (migrate_legacy + source_data.MICRO_SUFFIX).
"""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import source_data as SD

XLSX = os.path.join(ROOT, 'Банк_фраз.xlsx')
OUT  = os.path.join(ROOT, 'data', 'phrases.json')


def _find_header_row(rows, must=('ID',)):
    """Ищет строку заголовков среди первых 8 строк листа."""
    for i, r in enumerate(rows[:8]):
        cells = [str(c).strip() if c is not None else '' for c in r]
        if any(c in must for c in cells) and sum(1 for c in cells if c) >= 2:
            return i
    return None


def read_sheet(ws, must=('ID',)):
    """Возвращает список словарей {заголовок: значение}. Строка заголовков ищется автоматически."""
    rows = list(ws.iter_rows(values_only=True))
    hi = _find_header_row(rows, must)
    if hi is None:
        return []
    headers = [str(h).strip() if h is not None else '' for h in rows[hi]]
    out = []
    for r in rows[hi + 1:]:
        if r is None or all(v is None or str(v).strip() == '' for v in r):
            continue
        out.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
    return out


def yesno(v, default=False):
    if v is None:
        return default
    return str(v).strip().lower() in ('да', '1', 'true', 'yes', '+')


def idlist(v):
    if not v:
        return []
    return [x.strip() for x in re.split(r'[;,]', str(v)) if x.strip()]


def parse_key(v):
    """'as: ID1;ID2 | goal: GL-01' -> {'as': [...], 'goal': [...]}"""
    res = {}
    if not v:
        return res
    for part in str(v).split('|'):
        if ':' not in part:
            continue
        k, val = part.split(':', 1)
        k = k.strip().lower()
        res[k] = [x.strip() for x in re.split(r'[;,]', val) if x.strip()]
    for k in ('as', 'goal', 'meth', 'rec', 'res', 'nxt', 'svc'):
        res.setdefault(k, [])
    return res


def cell(d, *names, default=''):
    for n in names:
        if n in d and d[n] is not None:
            return d[n]
    return default


def migrate_legacy(data):
    """Старый Excel-банк хранит прежние коды (AS1-S3-01, микроблок AS4, RC1…).
    Переводим их на семантические коды (AS1sens-S3-01, AS4ini, RC1home…) по карте
    code_migration из мастер-данных; неизвестные ID старого формата (например, своя
    строка AS5-S2-99) переводятся тем же правилом через SD.rename_code.
    Новые коды и коды других блоков (GL-…, MT-…, SV-…) не меняются."""
    mig = data.get('code_migration') or {}

    def m(x):
        if not x:
            return x
        return mig.get(x) or SD.rename_code(x)

    for p in data.get('phrases', []):
        p['id'] = m(p.get('id'))
        p['micro'] = m(p.get('micro'))
        for k in ('links_goal', 'links_meth', 'links_rec'):
            p[k] = [m(x) for x in (p.get(k) or [])]
    for q in data.get('questionnaire', []) + data.get('survey_core', []):
        if q.get('micro'):
            q['micro'] = m(q['micro'])
        for kk in (q.get('key') or {}).values():
            for f in ('as', 'goal', 'meth', 'rec', 'res', 'nxt'):
                if kk.get(f):
                    kk[f] = [m(x) for x in kk[f]]
    for mb in data.get('microblocks', []):
        mb['id'] = m(mb.get('id'))
    for a in data.get('abilities', []):
        if a.get('micro'):
            a['micro'] = m(a['micro'])
    for s in data.get('service_phrases', []):
        if s.get('microblock'):
            s['microblock'] = m(s['microblock'])
    for sp in data.get('spheres', []):
        if sp.get('micro'):
            sp['micro'] = m(sp['micro'])
        if sp.get('bank_as'):
            sp['bank_as'] = [m(x) for x in sp['bank_as']]
    for sid, keys in (data.get('sphere_keys') or {}).items():
        for kk in keys.values():
            for f in ('as', 'goal', 'meth', 'rec', 'res', 'nxt'):
                if kk.get(f):
                    kk[f] = [m(x) for x in kk[f]]
    return data


def load(xlsx_path=XLSX):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    names = wb.sheetnames

    data = SD.build()   # базовая структура (meta, status_by_answer, doc_structure ...)

    if 'Фразы' in names:
        phrases = []
        for r in read_sheet(wb['Фразы']):
            pid = str(cell(r, 'ID')).strip()
            txt = str(cell(r, 'Текст формулировки')).strip()
            if not pid or not txt:
                continue
            phrases.append({
                'id': pid,
                'dir': str(cell(r, 'Направление', default='IV')).strip() or 'IV',
                'block': str(cell(r, 'Блок', default='AS')).strip(),
                'micro': str(cell(r, 'Микроблок')).strip(),
                'ability': str(cell(r, 'Способность')).strip(),
                'status': str(cell(r, 'Статус')).strip(),
                'text': txt,
                'source': str(cell(r, 'Источник')).strip(),
                'needs_example': yesno(cell(r, 'Требует примера'), True),
                'example_hint': str(cell(r, 'Подсказка: какой пример нужен')).strip(),
                'links_goal': idlist(cell(r, 'Цели(ID)')),
                'links_meth': idlist(cell(r, 'Приёмы(ID)')),
                'links_rec': idlist(cell(r, 'Рекомендации(ID)')),
                'tags': str(cell(r, 'Теги/домен')).strip(),
                'active': yesno(cell(r, 'Активна'), True),
            })
        if phrases:
            data['phrases'] = phrases

    if 'Опросник_ключ' in names:
        q = []
        for r in read_sheet(wb['Опросник_ключ']):
            qid = str(cell(r, 'ID')).strip()
            txt = str(cell(r, 'Пункт опросника')).strip()
            if not qid or not txt:
                continue
            stage = cell(r, 'Стадия DIR', default=0)
            try:
                stage = int(float(stage))
            except Exception:
                stage = 0
            q.append({
                'id': qid, 'dir': str(cell(r, 'Направление', default='IV')) or 'IV',
                'source': str(cell(r, 'Источник')) or 'DIR',
                'stage': stage,
                'domain': str(cell(r, 'Домен')).strip(),
                'micro': str(cell(r, 'Микроблок')).strip(),
                'text': txt, 'scale': True, 'note': '',
                'key': {
                    '3': parse_key(cell(r, 'Ключ: 3 (стабильно)')),
                    '2': parse_key(cell(r, 'Ключ: 2 (иногда)')),
                    '1': parse_key(cell(r, 'Ключ: 1 (с поддержкой)')),
                    '0': parse_key(cell(r, 'Ключ: 0 (не проявляется)')),
                    'N': parse_key(cell(r, 'Ключ: N (не проверялось)')),
                },
            })
        if q:
            data['questionnaire'] = q

    if 'Опросник_ядро' in names:
        core = []
        for r in read_sheet(wb['Опросник_ядро']):
            qid = str(cell(r, 'ID')).strip()
            txt = str(cell(r, 'Оценка (вопрос эксперта)')).strip()
            if not qid or not txt:
                continue
            core.append({
                'id': qid, 'dir': 'IV', 'source': 'ядро',
                'micro': str(cell(r, 'Микроблок')).strip(),
                'domain': str(cell(r, 'Область')).strip(),
                'text': txt, 'scale': True,
                'key': {
                    '3': parse_key(cell(r, 'Ключ: 3 (стабильно)')),
                    '2': parse_key(cell(r, 'Ключ: 2 (иногда)')),
                    '1': parse_key(cell(r, 'Ключ: 1 (с поддержкой)')),
                    '0': parse_key(cell(r, 'Ключ: 0 (не проявляется)')),
                    'N': parse_key(cell(r, 'Ключ: N (не проверялось)')),
                },
            })
        if core:
            data['survey_core'] = core

    if 'ИВ2_наблюдения' in names:
        iv2 = []
        for r in read_sheet(wb['ИВ2_наблюдения']):
            qid = str(cell(r, 'ID')).strip()
            txt = str(cell(r, 'Вопрос')).strip()
            if not qid or not txt:
                continue
            iv2.append({'id': qid, 'dir': 'IV', 'source': 'ИВ2',
                        'group': str(cell(r, 'Раздел')).strip(), 'text': txt,
                        'scale': False, 'micro': '', 'key': {}})
        if iv2:
            data['iv2_questions'] = iv2

    if 'Статусы' in names:
        st = []
        for r in read_sheet(wb['Статусы']):
            code = str(cell(r, 'Код')).strip()
            if not code:
                continue
            st.append({'code': code, 'name': str(cell(r, 'Название')).strip(),
                       'definition': str(cell(r, 'Когда ставим')).strip(),
                       'rule': str(cell(r, 'Правило')).strip(),
                       'goes_to': str(cell(r, 'Куда попадает')).strip()})
        if st:
            data['statuses'] = st

    if 'Поддержка' in names:
        sp = [{'id': re.sub(r'\W+', '', str(cell(r, 'Тип поддержки')).lower())[:8] or ('sup%d' % i),
               'name': str(cell(r, 'Тип поддержки')).strip(),
               'examples': str(cell(r, 'Примеры')).strip()}
              for i, r in enumerate(read_sheet(wb['Поддержка'])) if cell(r, 'Тип поддержки')]
        if sp:
            data['support_types'] = sp

    if 'Блоки' in names:
        bl = []
        for r in read_sheet(wb['Блоки']):
            bid = str(cell(r, 'ID')).strip()
            if not bid:
                continue
            try:
                order = int(float(cell(r, '№', default=0)))
            except Exception:
                order = len(bl) + 1
            bl.append({'id': bid, 'order': order, 'name': str(cell(r, 'Название блока')).strip(),
                       'mandatory': yesno(cell(r, 'Обязателен'), True),
                       'hint': str(cell(r, 'Что пишем')).strip()})
        if bl:
            data['blocks'] = sorted(bl, key=lambda x: x['order'])

    if 'Микроблоки' in names:
        mb = []
        for r in read_sheet(wb['Микроблоки']):
            mid = str(cell(r, 'ID')).strip()
            if not mid:
                continue
            try:
                order = int(float(cell(r, '№', default=0)))
            except Exception:
                order = len(mb) + 1
            svc = str(cell(r, 'Служебная фраза (ID)')).strip()
            mb_row = {'id': mid, 'block': str(cell(r, 'Блок')).strip(), 'order': order,
                      'name': str(cell(r, 'Название микротемы')).strip(),
                      'mandatory': yesno(cell(r, 'Обязателен'), True),
                      'service': svc or None}
            # микротемы форка v2 (сферы 8–9) не показываются в классической версии
            if yesno(cell(r, 'Только v2'), False):
                mb_row['v2only'] = True
            mb.append(mb_row)
        if mb:
            data['microblocks'] = mb

    if 'Способности' in names:
        ab = []
        for r in read_sheet(wb['Способности']):
            aid = str(cell(r, 'ID')).strip()
            if not aid:
                continue
            try:
                stage = int(float(cell(r, 'Стадия DIR', default=0)))
            except Exception:
                stage = 0
            ab.append({'id': aid, 'name': str(cell(r, 'Способность')).strip(),
                       'micro': str(cell(r, 'Микроблок')).strip(), 'stage': stage})
        if ab:
            data['abilities'] = ab

    if 'Служебные_фразы' in names:
        sv = []
        for r in read_sheet(wb['Служебные_фразы']):
            sid = str(cell(r, 'ID')).strip()
            if not sid:
                continue
            micro = str(cell(r, 'Микроблок')).strip()
            sv.append({'id': sid, 'microblock': (None if micro in ('', '—', 'None') else micro),
                       'situation': str(cell(r, 'Ситуация')).strip(),
                       'text': str(cell(r, 'Формулировка')).strip()})
        if sv:
            data['service_phrases'] = sv

    if 'Карта_связей' in names:
        lm = [{'marker': str(cell(r, 'Маркер / способность')).strip(),
               'stable': str(cell(r, 'Если стабильно')).strip(),
               'sometimes': str(cell(r, 'Если иногда')).strip(),
               'goal': str(cell(r, 'Цель')).strip(),
               'method': str(cell(r, 'Приём')).strip(),
               'rec': str(cell(r, 'Рекомендация')).strip()}
              for r in read_sheet(wb['Карта_связей']) if cell(r, 'Маркер / способность')]
        if lm:
            data['link_map'] = lm

    if 'Чек-лист' in names:
        ch = []
        for r in read_sheet(wb['Чек-лист'], must=('Проверка перед выдачей документа', '№')):
            txt = str(cell(r, 'Проверка перед выдачей документа')).strip()
            if not txt:
                continue
            imp = str(cell(r, 'Важность')).strip().lower()
            ch.append({'id': str(cell(r, '№')).strip() or ('C%d' % (len(ch) + 1)),
                       'text': txt,
                       'severity': {'обязательно': 'block', 'желательно': 'warn'}.get(imp, 'warn')})
        if ch:
            data['checklist'] = ch

    if 'Литература' in names:
        lit = [str(cell(r, 'Издание')).strip() for r in read_sheet(wb['Литература'])
               if str(cell(r, 'Издание')).strip()]
        if lit:
            data['literature'] = lit

    if 'Направления' in names:
        dr = []
        for r in read_sheet(wb['Направления']):
            did = str(cell(r, 'ID')).strip()
            if not did:
                continue
            dr.append({'id': did, 'name': str(cell(r, 'Название')).strip(),
                       'short': str(cell(r, 'Коротко')).strip(),
                       'subject_default': str(cell(r, 'Термин по умолчанию')).strip() or 'ребёнок',
                       'active': yesno(cell(r, 'Наполнено'), False),
                       'note': str(cell(r, 'Примечание')).strip()})
        if dr:
            data['directions'] = dr

    if 'Сферы' in names:
        sp = []
        for r in read_sheet(wb['Сферы']):
            sid = str(cell(r, 'ID')).strip()
            if not sid:
                continue
            try:
                order = int(float(cell(r, '№', default=0)))
            except Exception:
                order = len(sp) + 1
            sp_row = {'id': sid, 'order': order, 'name': str(cell(r, 'Название сферы')).strip(),
                      'short': str(cell(r, 'Коротко')).strip(),
                      'micro': str(cell(r, 'Микроблок документа')).strip()}
            fedl_blocks = idlist(cell(r, 'Блоки FEDL2i'))
            if fedl_blocks:
                sp_row['sens_blocks'] = fedl_blocks
            pool = idlist(cell(r, 'Пул фраз банка (ID)'))
            if pool:
                sp_row['bank_as'] = pool
            sp.append(sp_row)
        if sp:
            data['spheres'] = sp
    if 'Шкала_МКФ' in names:
        ps = []
        for r in read_sheet(wb['Шкала_МКФ'], must=('Оценка педагога',)):
            try:
                ped = int(float(cell(r, 'Оценка педагога')))
                mkf = int(float(cell(r, 'Балл МКФ')))
            except Exception:
                continue
            row = {'ped': ped, 'mkf': mkf, 'code': str(cell(r, 'Код')).strip(),
                   'name': str(cell(r, 'Название')).strip(),
                   'text': str(cell(r, 'Что означает (педагогическим языком)')).strip()}
            sup = str(cell(r, 'Критерий через поддержку (язык V1)',
                           'Критерий через поддержку')).strip()
            if sup:
                row['support'] = sup
            v1 = str(cell(r, 'Соответствие в V1')).strip()
            if v1:
                row['v1'] = v1
            ps.append(row)
        if ps:
            data['ped_scale'] = ps
    if 'Сигнальные_отметки' in names:
        sg = []
        for r in read_sheet(wb['Сигнальные_отметки']):
            sid = str(cell(r, 'ID')).strip()
            if not sid:
                continue
            svc = str(cell(r, 'Служебная фраза (ID)')).strip()
            row = {'id': sid, 'name': str(cell(r, 'Название')).strip(),
                   'text': str(cell(r, 'Значение')).strip(), 'svc': svc or None}
            v1 = str(cell(r, 'Соответствие в V1')).strip()
            if v1:
                row['v1'] = v1
            sg.append(row)
        if sg:
            data['ped_signals'] = sg
    if 'Ключ_сфер' in names:
        sk = {}
        for r in read_sheet(wb['Ключ_сфер']):
            sid = str(cell(r, 'Сфера(ID)')).strip()
            ped = str(cell(r, 'Оценка')).strip()
            if not sid or not ped:
                continue
            key = {'as': idlist(cell(r, 'Актуальное(ID)')), 'goal': idlist(cell(r, 'Цели(ID)')),
                   'meth': idlist(cell(r, 'Приёмы(ID)')), 'rec': idlist(cell(r, 'Рекомендации(ID)')),
                   'res': [], 'nxt': [], 'svc': idlist(cell(r, 'Служебные(ID)'))}
            try:
                pedk = int(float(ped))
            except Exception:
                pedk = ped
            sk.setdefault(sid, {})[pedk] = key
        if sk:
            data['sphere_keys'] = sk
    if 'Литература_правила' in names:
        lr = {}
        for r in read_sheet(wb['Литература_правила']):
            sid = str(cell(r, 'Сфера(ID)')).strip()
            if sid:
                lr[sid] = idlist(cell(r, 'Подстроки названий из списка «Литература»'))
        if lr:
            data['lit_rules'] = lr

    if 'Причины_S0' in names:
        pr = [str(cell(r, 'Причина невозможности оценки')).strip()
              for r in read_sheet(wb['Причины_S0'], must=('Причина невозможности оценки',))
              if str(cell(r, 'Причина невозможности оценки')).strip()]
        if pr:
            data['ped_reasons'] = pr

    return migrate_legacy(data)


def main():
    if not os.path.exists(XLSX):
        raise SystemExit('Файл не найден: ' + XLSX)
    data = load()
    errs = SD.validate(data)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print('Из Excel собрано: %d фраз, %d пунктов опросника, %d вопросов ИВ2'
          % (len(data['phrases']), len(data['questionnaire']), len(data['iv2_questions'])))
    print('Записано:', os.path.relpath(OUT, ROOT))
    if errs:
        print('\n! Предупреждения целостности (проверьте ID в связях и ключе):')
        for e in errs:
            print('   -', e)
    else:
        print('Целостность данных: OK')
    print('\nДальше выполните:  python tools/build.py')


if __name__ == '__main__':
    main()
